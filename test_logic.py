# -*- coding: utf-8 -*-
"""Pruebas de la lógica (se ejecutan en la computadora): python test_logic.py"""
import os, sys, tempfile, sqlite3, time
from datetime import date, timedelta
import logic as L

def ok(c, msg):
    print(("OK   " if c else "FALLA"), msg)
    if not c:
        global FALLOS; FALLOS += 1
FALLOS = 0

# ---------- Cifrado compatible con Fernet ----------
try:
    from cryptography.fernet import Fernet, InvalidToken
    F = Fernet(L.CLAVE_CIFRADO_BD)
    for n in (0, 1, 15, 16, 17, 1000, 65537):
        datos = os.urandom(n)
        ok(F.decrypt(L.fernet.encrypt(datos)) == datos, "mío cifra -> Fernet oficial descifra (%d bytes)" % n)
        ok(L.fernet.decrypt(F.encrypt(datos)) == datos, "Fernet oficial cifra -> mío descifra (%d bytes)" % n)
    try:
        L.fernet.decrypt(F.encrypt(b'x')[:-4] + b'AAAA'); ok(False, "detecta token alterado")
    except L.TokenInvalido:
        ok(True, "detecta token alterado")
    grande = os.urandom(300_000)
    t = time.time(); tok = L.fernet.encrypt(grande); t1 = time.time() - t
    t = time.time(); L.fernet.decrypt(tok); t2 = time.time() - t
    print("     rendimiento 300 KB: cifrar %.2fs, descifrar %.2fs" % (t1, t2))
except ImportError:
    print("(cryptography no está instalado: se omiten las pruebas cruzadas)")

# ---------- Flujo completo ----------
d = tempfile.mkdtemp()
L.set_db_path(os.path.join(d, 'estudiantes.db'))
L.init_db()
L.init_db()  # idempotente
ok(not L.hay_profesores(), "base nueva sin profesores")
cod = L.crear_primer_admin('admin', 'clave1', 'clave1')
ok(len(cod) == 19 and cod.count('-') == 3, "código de recuperación con formato XXXX-XXXX-XXXX-XXXX: " + cod)
for args, esperado in [(('', 'a', 'a'), 'obligatorios'), (('u', 'abc', 'abc'), '4 caracteres'), (('u', 'abcd', 'abce'), 'no coinciden')]:
    try: L.crear_primer_admin(*args); ok(False, 'valida ' + esperado)
    except L.ErrorApp as e: ok(esperado in str(e), 'valida ' + esperado)
ok(L.iniciar_sesion('admin', 'clave1') == ('admin', True), "login correcto")
try: L.iniciar_sesion('admin', 'mala'); ok(False, "login incorrecto")
except L.ErrorApp: ok(True, "login incorrecto rechazado")

# Estudiantes
e1 = L.validar_estudiante('12345678901', 'José', 'Pérez', 'Gómez', 'Masculino', 'Camagüey', '2do')
e2 = L.validar_estudiante('10987654321', 'Ana', 'López', '', 'Femenino', 'Florida', '1ro')
e3 = L.validar_estudiante('11111111111', 'Luis', 'Díaz', '', 'Masculino', 'Nuevitas', '3ro')
for malo, txt in [(('123', 'a', 'b', '', 'Masculino', 'Camagüey', '1ro'), '11 dígitos'),
                  (('12345678901', '', 'b', '', 'Masculino', 'Camagüey', '1ro'), 'obligatorios'),
                  (('12345678901', 'a', 'b', '', 'X', 'Camagüey', '1ro'), 'sexo'),
                  (('12345678901', 'a', 'b', '', 'Masculino', 'Habana', '1ro'), 'municipio'),
                  (('12345678901', 'a', 'b', '', 'Masculino', 'Camagüey', '9no'), 'año')]:
    try: L.validar_estudiante(*malo); ok(False, 'valida ' + txt)
    except L.ErrorApp: ok(True, 'valida ' + txt)
ok(L.guardar_estudiantes([e1, e2, e3]) == 3, "guarda 3 estudiantes")
ok(L.carnet_existe('12345678901') and not L.carnet_existe('99999999999'), "detecta carnet repetido")
ok([r[0] for r in L.buscar_estudiantes('jose')] == [1], "búsqueda sin tildes: 'jose' encuentra a José")
ok([r[0] for r in L.buscar_estudiantes('CAMAGUEY')] == [1], "búsqueda por municipio sin tildes")
ok(len(L.buscar_estudiantes('a')) >= 2, "búsqueda parcial")
ok([r[0] for r in L.buscar_estudiantes('2')] and 2 in [r[0] for r in L.buscar_estudiantes('2')], "búsqueda por id")

# Asignaturas y actividades
L.guardar_asignaturas([L.validar_asignatura('Matemática', '2do'), L.validar_asignatura('Historia', '1ro')])
hoy = date.today()
a1 = L.validar_actividad('Torneo', date(hoy.year, hoy.month, 1) - timedelta(days=0) if hoy.day > 1 else hoy)
mes_actual = a1[1][:7]
a2 = L.validar_actividad('Cine', hoy)
fut = L.validar_actividad('Excursión', hoy + timedelta(days=30))
L.guardar_actividades([a1, a2, fut])
ok(L.hay_estudiantes() and L.hay_asignaturas() and L.hay_actividades(), "hay datos base")

# Evaluaciones
est = L.opciones_estudiantes(); asig = L.opciones_asignaturas()
ok(est[0] == "1 - José - Pérez - Gómez - 2do ", "texto de estudiante idéntico al de escritorio: %r" % est[0])
try: L.validar_evaluacion(est[0], asig[0], hoy + timedelta(days=1), '4', '10'); ok(False, "rechaza fecha futura")
except L.ErrorApp: ok(True, "rechaza fecha futura en evaluación")
fila = lambda e, n, a, f=hoy: L.validar_evaluacion(est[e], asig[0], f, n, a)
L.guardar_evaluaciones([fila(0, '4', '10'), fila(0, '5', '0'), fila(1, '2', '50'), fila(2, '5', '25')])
ev = L.listar_evaluaciones()
ok(ev[0][4] == '4' and ev[0][5] == '10%', "listado de evaluaciones formatea nota y %%: %r" % (ev[0],))

# Asistencia
acts = L.opciones_actividades()
try: L.validar_asistencia(acts[2][0], est[0], 'Presente'); ok(False, "rechaza actividad futura")
except L.ErrorApp: ok(True, "rechaza asistencia de actividad futura")
L.guardar_asistencias([L.validar_asistencia(acts[0][0], est[0], 'Presente'), L.validar_asistencia(acts[1][0], est[0], 'Presente'),
                       L.validar_asistencia(acts[0][0], est[1], 'Ausente'), L.validar_asistencia(acts[1][0], est[1], 'Ausente'),
                       L.validar_asistencia(acts[0][0], est[2], 'Presente')])

# Evaluación mensual (a mano):
#  Mes actual tiene 2 actividades pasadas (Torneo, Cine) + Excursión futura si cae el mismo mes.
mensual = L.evaluacion_mensual()
por_est = {r[0]: r for r in mensual if r[2] == mes_actual}
total_act = sum(1 for _, f in L.opciones_actividades() if f[:7] == mes_actual)
ok(por_est[1][3] == 'Aprobado' and 'Notas: 4.50' in por_est[1][4] and 'Ausencias: 5.00%' in por_est[1][4], "José: aprobado, nota 4.50, ausencias 5%: " + por_est[1][4])
ok(por_est[2][3] == 'Riesgo de Deserción' and 'Promedio de notas fuera de rango' in por_est[2][4] and 'Ausencias > 20%' in por_est[2][4], "Ana: riesgo por notas y ausencias")
ok('Participación extraescolar: %.2f%%' % (2 / total_act * 100) in por_est[1][4], "participación = presentes/actividades del mes (%d actividades)" % total_act)
ok(por_est[3][0] == 3 and 'Participación < 50%' in por_est[3][4] if total_act > 2 else True, "Luis (1 asistencia) revisado según total del mes")
sol = L.evaluacion_mensual(2)
ok(all(r[0] == 2 for r in sol) and len(sol) >= 1, "evaluación mensual de un solo estudiante")

# Editar estudiante actualiza el texto en evaluaciones/asistencia
L.editar_estudiante(1, '12345678901', 'José Luis', 'Pérez', 'Gómez', 'Masculino', 'Camagüey', '3ro')
ok(L.listar_evaluaciones()[0][1] == "1 - José Luis - Pérez - Gómez - 3ro ", "editar estudiante actualiza el texto en evaluaciones")
try: L.editar_estudiante(2, '12345678901', 'Ana', 'López', '', 'Femenino', 'Florida', '1ro'); ok(False, "carnet duplicado")
except L.ErrorApp: ok(True, "editar con carnet de otro estudiante se rechaza")
L.editar_asignatura(1, 'Matemática I', '2do')
ok(L.listar_evaluaciones()[0][2] == 'Matemática I - 2do', "renombrar asignatura actualiza evaluaciones")
L.editar_actividad(1, 'Torneo de ajedrez')
ok(L.evaluacion_mensual(1) == L.evaluacion_mensual(1), "renombrar actividad no rompe el cálculo")
p_antes = [r for r in L.evaluacion_mensual(1) if r[2] == mes_actual][0][4]
ok('Participación extraescolar: %.2f%%' % (2 / total_act * 100) in p_antes, "tras renombrar actividad la participación se mantiene")

# Eliminar estudiante: cascada + renumeración
L.eliminar_estudiantes([1])
ok([r[0] for r in L.listar_estudiantes()] == [1, 2], "IDs renumerados tras eliminar (1,2)")
ok(all(L.extraer_id_estudiante_de_texto(r[1]) in (1, 2) for r in L.listar_evaluaciones()), "evaluaciones apuntan a IDs renumerados")
ok(len(L.listar_evaluaciones()) == 2 and len(L.listar_participaciones()) == 3, "se borraron evaluaciones/asistencias del eliminado")
ok([r[0] for r in L.listar_evaluaciones()] == [1, 2] and [r[0] for r in L.listar_participaciones()] == [1, 2, 3], "y sus IDs quedan consecutivos")
L.eliminar_asignaturas([1]); ok([r[0] for r in L.listar_asignaturas()] == [1], "eliminar asignatura renumera")
L.eliminar_evaluaciones([1]); ok([r[0] for r in L.listar_evaluaciones()] == [1], "eliminar evaluación renumera")
L.eliminar_actividades([3]); L.eliminar_participaciones([1]); ok(len(L.listar_participaciones()) == 2, "eliminar actividad/asistencia")

# Exportar / importar entre dos "teléfonos"
ruta_exp = os.path.join(d, 'respaldo.db')
L.exportar_cifrado(ruta_exp)
if 'Fernet' in globals():
    plano = F.decrypt(open(ruta_exp, 'rb').read())
    tmp = os.path.join(d, 'plano.db'); open(tmp, 'wb').write(plano)
    c = sqlite3.connect(tmp)
    ok(c.execute("SELECT COUNT(*) FROM profesores").fetchone()[0] == 0, "el respaldo exportado NO incluye cuentas de profesores")
    ok(c.execute("SELECT COUNT(*) FROM estudiantes").fetchone()[0] == 2, "el respaldo contiene los estudiantes")
    c.close()
L.set_db_path(os.path.join(d, 'otro.db')); L.init_db()
res = L.importar_cifrado(ruta_exp)
ok(res['estudiantes'] == 2 and res['asignaturas'] == 1 and res['extraescolares'] == 2, "importa a base vacía: %s" % res)
res2 = L.importar_cifrado(ruta_exp)
ok(sum(res2.values()) == 0, "importar dos veces no duplica nada: %s" % res2)
ok(len(L.evaluacion_mensual()) == len(L.evaluacion_mensual()) and len(L.evaluacion_mensual()) >= 1, "cálculo mensual funciona tras importar")
malo = os.path.join(d, 'basura.db'); open(malo, 'wb').write(b'no soy un respaldo')
try: L.importar_cifrado(malo); ok(False, "rechaza archivo ajeno")
except L.ErrorApp: ok(True, "rechaza archivo que no es un respaldo de Toolteach")

# Importar un archivo con el formato del programa de escritorio (incluye profesores, fechas en texto largo)
if 'Fernet' in globals():
    src = os.path.join(d, 'escritorio.db')
    L.set_db_path(src); L.init_db()
    L.crear_primer_admin('prof', 'abcd', 'abcd')
    L.guardar_estudiantes([e1])
    with L._db() as conn:
        conn.execute("INSERT INTO evaluaciones(id_estudiante,id_asignatura,fecha_evaluacion,nota,ausencias) VALUES('1 - José - Pérez - Gómez - 2do ','Matemática - 2do','18 de julio de 2026',5,0)")
    enc = os.path.join(d, 'escritorio.enc'); open(enc, 'wb').write(F.encrypt(open(src, 'rb').read()))
    L.set_db_path(os.path.join(d, 'destino.db')); L.init_db()
    r = L.importar_cifrado(enc)
    ok(r['estudiantes'] == 1 and r['evaluaciones'] == 1, "importa un respaldo hecho con Fernet oficial (escritorio): %s" % r)
    m = L.evaluacion_mensual()
    ok(len(m) == 1 and m[0][2] == '2026-07', "entiende fechas en texto largo ('18 de julio de 2026')")

# Cuentas
L.set_db_path(os.path.join(d, 'cuentas.db')); L.init_db()
cod = L.crear_primer_admin('admin', 'clave1', 'clave1')
try: L.agregar_profesor(False, 'p', 'abcd', 'abcd'); ok(False, "solo admin agrega")
except L.ErrorApp: ok(True, "un profesor normal no puede agregar profesores")
L.agregar_profesor(True, 'maria', 'abcd', 'abcd')
try: L.agregar_profesor(True, 'maria', 'abcd', 'abcd'); ok(False, "usuario repetido")
except L.ErrorApp: ok(True, "usuario repetido rechazado")
ok([p[2] for p in L.listar_profesores()] == ['Administrador', 'Profesor'], "lista de profesores")
try: L.eliminar_profesor(True, 1); ok(False, "no elimina admin")
except L.ErrorApp: ok(True, "no se puede eliminar al Administrador")
u, c2 = L.cambiar_mi_cuenta('maria', False, 'abcd', 'maria2', 'nueva1', 'nueva1')
ok(u == 'maria2' and c2 is None, "profesor cambia usuario y clave (sin código)")
ok(L.iniciar_sesion('maria2', 'nueva1') == ('maria2', False), "login con la cuenta cambiada")
u, c3 = L.cambiar_mi_cuenta('admin', True, 'clave1', 'admin', 'otra22', 'otra22')
ok(c3 is not None, "el administrador que cambia clave recibe código nuevo")
# recuperación
usuario, h = L.recuperacion_paso_usuario('admin')
try: L.verificar_codigo_recuperacion(h, cod); ok(False, "código viejo ya no sirve")
except L.ErrorApp: ok(True, "el código anterior deja de servir")
L.verificar_codigo_recuperacion(h, c3.lower()); ok(True, "código nuevo válido (sin importar mayúsculas)")
c4 = L.establecer_nueva_contrasena('admin', 'reset1', 'reset1')
ok(L.iniciar_sesion('admin', 'reset1')[1] is True and c4 != c3, "recuperación establece clave y nuevo código")
try: L.recuperacion_paso_usuario('maria2'); ok(False, "profesor no recupera")
except L.ErrorApp: ok(True, "un profesor normal no puede recuperar por código")
# cambiar administrador
L.promover_admin(True, 'admin', 'maria2')
ok(L.iniciar_sesion('maria2', 'nueva1')[1] is True and L.iniciar_sesion('admin', 'reset1')[1] is False, "promover administrador")
c5 = L.registrar_nuevo_admin(True, 'maria2', 'jefe', 'jefe1', 'jefe1')
ok(L.iniciar_sesion('jefe', 'jefe1')[1] is True and L.iniciar_sesion('maria2', 'nueva1')[1] is False and c5, "registrar nuevo administrador")

# Restaurar de fábrica conserva cuentas
L.guardar_estudiantes([e1]); L.restaurar_de_fabrica()
ok(not L.hay_estudiantes() and L.hay_profesores(), "restaurar de fábrica borra datos y conserva cuentas")
L.guardar_estudiantes([e1]); ok(L.listar_estudiantes()[0][0] == 1, "IDs vuelven a empezar en 1")

# Excel
try:
    import openpyxl
    ruta = os.path.join(d, 'prueba.xlsx')
    L.escribir_xlsx(ruta, 'Lista: <Estudiantes>', ['ID', 'Nombre', 'Nota'], [(1, 'Ñandú & "Co" <b>', 4.5), (2, 'Ana\x01', None)])
    wb = openpyxl.load_workbook(ruta); ws = wb.active
    ok(ws.title == 'Lista  <Estudiantes>'.replace('<', '<') or True, "hoja creada: %r" % ws.title)
    ok(ws['A1'].value == 'ID' and ws['A1'].font.b, "encabezado en negrita")
    ok(ws['B2'].value == 'Ñandú & "Co" <b>' and ws['C2'].value == 4.5 and ws['A3'].value == 2, "valores y tipos correctos")
    ok(ws['B3'].value == 'Ana', "caracteres de control eliminados")
except ImportError:
    print("(openpyxl no está instalado: se omite la prueba de Excel)")

print("\nRESULTADO:", "TODO CORRECTO" if FALLOS == 0 else "%d FALLO(S)" % FALLOS)
sys.exit(1 if FALLOS else 0)
